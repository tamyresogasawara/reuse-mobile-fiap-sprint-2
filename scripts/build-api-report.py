from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "TamyresOgasawaraMacedo-ReUse-APIs-Auth-Cache.docx"
REPO_URL = "https://github.com/tamyresogasawara/reuse-mobile-fiap-sprint-2"
PREVIEW_URL = "https://tamyresogasawara.github.io/reuse-mobile-fiap-sprint-2/"
ASSETS = ROOT / "docs" / "assets"

GREEN = "173F2A"
GREEN_SOFT = "42634F"
LIME = "D7FF78"
LIME_SOFT = "F4FFD7"
INK = "101713"
MUTED = "5E6962"
BORDER = "CED8D0"
ORANGE = "E8783D"
WHITE = "FFFFFF"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str, color=GREEN_SOFT):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color_node, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([props, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_title(doc: Document, text: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)
    run.font.size = Pt(23)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(12)
        r2 = p2.add_run(subtitle)
        r2.font.color.rgb = RGBColor.from_string(GREEN_SOFT)
        r2.font.size = Pt(11)
        r2.bold = True


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_tag(doc: Document, text: str, fill=LIME_SOFT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, 70, 140, 70, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(GREEN)
    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(MUTED)
    r.font.size = Pt(9)


def add_picture(doc: Document, filename: str, width_cm: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
    add_caption(doc, caption)


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    for name, size, color in (("Title", 34, INK), ("Heading 1", 18, GREEN), ("Heading 2", 13, GREEN_SOFT)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "REUSE  ·  FIAP  ·  MOBILE DEVELOPMENT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(GREEN_SOFT)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]], widths: list) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, title, width in zip(table.rows[0].cells, headers, widths):
        cell.width = width
        shade(cell, GREEN)
        set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for idx, values in enumerate(rows):
        for cell, value, width in zip(table.add_row().cells, values, widths):
            cell.width = width
            shade(cell, LIME_SOFT if idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell)
            cell.paragraphs[0].add_run(value)
    doc.add_paragraph()


def add_picture_row(doc: Document, pictures: list[tuple[str, str]], width_cm: float) -> None:
    table = doc.add_table(rows=2, cols=len(pictures))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, (filename, caption) in enumerate(pictures):
        picture = table.cell(0, col).paragraphs[0]
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
        text = table.cell(1, col).paragraphs[0]
        text.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = text.add_run(caption)
        run.italic = True
        run.font.color.rgb = RGBColor.from_string(MUTED)
        run.font.size = Pt(9)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    configure_document(doc)

    # Capa
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(55)
    add_tag(doc, "ATIVIDADE · APIs, AUTENTICAÇÃO E CACHE LOCAL", LIME)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("ReUse Mobile")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(38)
    r.font.color.rgb = RGBColor.from_string(INK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(22)
    r2 = p2.add_run("Integração com APIs, sessões persistentes e experiência offline")
    r2.font.size = Pt(18)
    r2.font.color.rgb = RGBColor.from_string(GREEN_SOFT)
    r2.bold = True

    cover_table = doc.add_table(rows=4, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    labels = [
        ("Aluna", "Tamyres Ogasawara Macedo"),
        ("RM", "552055"),
        ("Modalidade", "Trabalho individual"),
        ("Instituição", "FIAP"),
    ]
    for row, (label, value) in zip(cover_table.rows, labels):
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(12)
        for cell in row.cells:
            set_cell_margins(cell, 120, 150, 120, 150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(row.cells[0], GREEN)
        shade(row.cells[1], "FFFFFF")
        a = row.cells[0].paragraphs[0].add_run(label)
        a.bold = True
        a.font.color.rgb = RGBColor.from_string(WHITE)
        b = row.cells[1].paragraphs[0].add_run(value)
        b.bold = True
        b.font.color.rgb = RGBColor.from_string(INK)

    doc.add_paragraph()
    repo = doc.add_paragraph()
    repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repo.add_run("Repositório GitHub: ").bold = True
    add_hyperlink(repo, REPO_URL.removeprefix("https://"), REPO_URL)
    published = doc.add_paragraph()
    published.alignment = WD_ALIGN_PARAGRAPH.CENTER
    published.add_run("Aplicação web: ").bold = True
    add_hyperlink(published, PREVIEW_URL.removeprefix("https://").rstrip("/"), PREVIEW_URL)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.add_run("Outubro de 2026").font.color.rgb = RGBColor.from_string(MUTED)

    # Resumo e arquitetura
    add_title(doc, "1. Visão geral da solução", "Como a implementação atende aos três critérios da atividade")
    add_body(doc, "A ReUse da Sprint 2 — com busca, favoritos, anúncios por câmera e gamificação — passou a ser conectada à API pública DummyJSON. O fluxo agora vai das boas-vindas ao login, restaura a sessão ao reabrir o aplicativo e oferece um catálogo online com cache local e fallback offline, sem perder nenhuma funcionalidade anterior.")
    add_table(doc, ["Critério", "Peso", "Entrega implementada"], [
        ("Consumo de API", "60%", "Login por POST e catálogo por GET, com validação da resposta, carregamento, atualização manual, retry e mensagens de erro."),
        ("Autenticação e sessões", "20%", "Tokens e perfil persistidos; restauração automática; nome e usuário na Início e no Perfil; logout."),
        ("Cache local", "20%", "Catálogo salvo por 15 minutos; cache recente sem rede e último cache como fallback offline."),
    ], [Cm(4), Cm(2.2), Cm(10.4)])

    add_heading(doc, "Arquitetura principal", 1)
    add_bullet(doc, "SessionProvider (src/state/SessionContext.tsx): restaura e valida a sessão antes de montar as rotas e expõe login e logout.")
    add_bullet(doc, "App.tsx: sem sessão, mostra Boas-vindas e Login; com sessão, libera as abas, detalhes, catálogo online, Meus anúncios e Meu impacto.")
    add_bullet(doc, "src/services/auth.ts: autenticação, normalização, persistência, restauração, expiração e logout.")
    add_bullet(doc, "src/services/products.ts: consumo do catálogo, validação, TTL e fallback do cache.")
    add_bullet(doc, "ExploreScreen (Catálogo online): estados de rede, cache, offline, erro, retry e atualização manual.")
    add_bullet(doc, "Testes automatizados: serviços, telas, navegação, sessão, gamificação e movimento.")

    # APIs
    add_title(doc, "2. Consumo de APIs", "APIs utilizadas, telas atualizadas e código-fonte")
    add_heading(doc, "2.1 Autenticação", 1)
    add_body(doc, "Endpoint: POST https://dummyjson.com/auth/login")
    add_body(doc, "A LoginScreen valida os campos e envia username, password e validade de 60 minutos. A resposta é validada antes de gerar uma sessão com accessToken, refreshToken e dados normalizados do usuário. Credenciais recusadas (HTTP 400) e outras falhas produzem mensagens claras sem liberar a navegação.")
    add_bullet(doc, "Dados utilizados: accessToken, refreshToken, id, username, firstName e lastName.")
    add_bullet(doc, "Conta pública de avaliação: emilys / emilyspass.")

    add_heading(doc, "2.2 Catálogo de produtos", 1)
    add_body(doc, "Endpoint: GET https://dummyjson.com/products?limit=12&select=id,title,description,price,category,thumbnail")
    add_body(doc, "O cartão Catálogo online da Início abre a ExploreScreen, que apresenta doze produtos recebidos da API com miniatura, título, categoria, descrição e preço formatado em reais. A tela mostra a origem dos dados: rede, cache recente ou cache expirado em modo offline.")
    add_bullet(doc, "Estados visuais: carregando, sucesso, erro, retry, cache e offline; Atualizar catálogo força nova requisição.")

    add_heading(doc, "2.3 Telas atualizadas", 1)
    add_table(doc, ["Tela", "O que mudou"], [
        ("Boas-vindas (WelcomeScreen)", "O botão Começar a explorar passa a levar ao login."),
        ("Login (LoginScreen) — nova", "Formulário integrado ao POST /auth/login, com validação e mensagens de erro."),
        ("Início (HomeScreen)", "Saudação com o nome da sessão e cartão que abre o Catálogo online."),
        ("Catálogo online (ExploreScreen) — nova", "Produtos do GET /products, com cache local, aviso offline, retry e atualização manual."),
        ("Perfil (ProfileScreen)", "Nome e usuário da sessão e botão Sair."),
    ], [Cm(5.6), Cm(11)])

    add_heading(doc, "2.4 Código-fonte React Native", 1)
    p = doc.add_paragraph()
    p.add_run("GitHub: ").bold = True
    add_hyperlink(p, REPO_URL, REPO_URL)
    p = doc.add_paragraph()
    p.add_run("Aplicação web: ").bold = True
    add_hyperlink(p, PREVIEW_URL, PREVIEW_URL)

    # Sessão
    add_title(doc, "3. Autenticação e sessões", "Telas em que foram aplicadas")
    add_table(doc, ["Tela", "Uso da autenticação e da sessão"], [
        ("Boas-vindas e Login", "Exibidas somente sem sessão válida; o login cria e salva a sessão."),
        ("Início", "Mostra o nome do usuário autenticado."),
        ("Perfil", "Mostra nome e usuário; o botão Sair encerra a sessão."),
        ("Demais telas", "Abas, Detalhes, Catálogo online, Meus anúncios e Meu impacto só são montadas com sessão válida."),
    ], [Cm(4.6), Cm(12)])
    add_picture_row(doc, [
        ("s4-01-login.png", "Figura 1 — Login integrado à API DummyJSON."),
        ("s2-02-inicio.png", "Figura 2 — Início com o nome da sessão."),
        ("s2-08-perfil.png", "Figura 3 — Perfil com usuário e opção Sair."),
    ], 4.4)
    add_body(doc, "No Android e iOS, os tokens são persistidos com Expo SecureStore; na web, com Async Storage, porque o SecureStore não existe no navegador. Ao iniciar, o SessionProvider restaura e valida a sessão salva: se for válida, o app abre direto na navegação principal; se tiver expirado (60 minutos) ou estiver corrompida, é descartada e a pessoa volta às boas-vindas.")

    # Cache
    add_title(doc, "4. Caching local", "Itens em cache e telas em que foi aplicado")
    add_table(doc, ["Item em cache", "Conteúdo", "Validade", "Telas"], [
        ("Catálogo de produtos (@reuse/products-v1)", "12 produtos da API (id, title, description, price, category, thumbnail) e savedAt", "15 minutos; depois vira fallback offline", "Catálogo online"),
        ("Sessão (reuse-session-v1)", "accessToken, refreshToken, expiresAt, id, username e nome", "60 minutos", "Login, Início e Perfil; restaurada ao abrir o app"),
    ], [Cm(4.2), Cm(5.8), Cm(3.4), Cm(3.2)])
    add_body(doc, "O catálogo fica no Async Storage. A sessão fica no Expo SecureStore no Android/iOS e no Async Storage na web. Favoritos, rascunho e anúncios da sprint anterior continuam salvos localmente no Async Storage.")
    add_picture_row(doc, [("s4-02-catalogo-api.png", "Figura 4 — Catálogo com dados atualizados pela API."), ("s4-03-catalogo-offline.png", "Figura 5 — Rede indisponível (simulada): último cache com aviso.")], 4.0)
    add_heading(doc, "Regras de leitura", 1)
    add_bullet(doc, "Cache recente: retorna imediatamente sem chamada à API.")
    add_bullet(doc, "Cache ausente ou expirado: consulta a rede e salva uma nova versão.")
    add_bullet(doc, "Falha de rede com cache disponível: apresenta o último catálogo com aviso offline.")
    add_bullet(doc, "Falha sem cache: mostra erro e botão para tentar novamente. Falhas do armazenamento não impedem o uso dos dados da rede.")

    # Evidências
    add_title(doc, "5. Testes e verificação", "Evidências executadas sobre a implementação final")
    add_table(doc, ["Verificação", "Resultado"], [
        ("Jest", "38 testes aprovados em 6 suítes"),
        ("TypeScript", "typecheck sem erros"),
        ("Build Expo Web", "exportação concluída e publicada no GitHub Pages"),
        ("Integração ao vivo", "login real com emilys e 12 produtos retornados pela DummyJSON"),
        ("Playwright", "build web de produção com a API real: login, sessão, catálogo e cache; modo offline simulado bloqueando a requisição do catálogo"),
        ("GitHub", "repositório tamyresogasawara/reuse-mobile-fiap-sprint-2"),
    ], [Cm(6), Cm(10.6)])

    add_heading(doc, "Cenários automatizados", 1)
    add_bullet(doc, "Autenticação externa e normalização da resposta.")
    add_bullet(doc, "Persistência, restauração, expiração e encerramento da sessão.")
    add_bullet(doc, "Validação do formulário e credenciais recusadas pela API.")
    add_bullet(doc, "Catálogo obtido pela API e salvo no cache; cache recente sem chamada de rede.")
    add_bullet(doc, "Fallback do cache expirado; falhas de leitura e gravação do cache.")
    add_bullet(doc, "Erro sem cache com retry; navegação autenticada, catálogo online e logout.")
    add_bullet(doc, "Fluxos anteriores preservados: busca, favoritos, anúncios, perfil e gamificação.")

    # Conclusão
    add_title(doc, "6. Conclusão", "Resultado e aderência ao objetivo da atividade")
    add_body(doc, "A ReUse atende aos três eixos de avaliação. A API sustenta tanto o acesso quanto o catálogo online. A sessão persistente protege o fluxo principal e melhora a continuidade de uso. O cache local reduz chamadas desnecessárias e mantém o catálogo disponível durante falhas temporárias de conexão.")
    add_body(doc, "A integração foi feita sobre a base da Sprint 2, preservando a identidade visual, as animações, a gamificação e os fluxos locais, com separação entre telas, serviços, estado e testes.")
    add_heading(doc, "Links finais", 1)
    p = doc.add_paragraph()
    p.add_run("Código-fonte: ").bold = True
    add_hyperlink(p, REPO_URL, REPO_URL)
    p2 = doc.add_paragraph()
    p2.add_run("Aplicação web: ").bold = True
    add_hyperlink(p2, PREVIEW_URL, PREVIEW_URL)
    add_tag(doc, "ENTREGA INDIVIDUAL · TAMYRES OGASAWARA MACEDO", ORANGE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
